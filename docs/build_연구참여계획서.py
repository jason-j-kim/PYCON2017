#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""연구 참여 계획서 — KDI 가 수행할 연구에 한 명의 연구자로 참여하겠다는 문서.

기관의 사업계획도, 단독 연구계획도 아니다. 주관은 KDI 이고 본인은 참여자다.
그래서 문서가 답해야 할 질문은 하나다 — 이 사람을 넣으면 무엇이 달라지는가.

세 역할로 답한다. 설계 자문, 구현, 그리고 아이디어 내생 성장론 관점의
보고서. 앞의 둘은 이미 만들어 둔 것이 근거가 되고, 셋째는 이 연구가
스스로 갖기 어려운 이론 축을 더하는 일이다.

셋째에 지면을 더 준다. 앞의 둘은 증명이 끝난 것이고 셋째는 제안이기 때문이다.
"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

KO = "맑은 고딕"
TITLE, HEAD, BODY, SMALL = Pt(15), Pt(12), Pt(11), Pt(10)
BLACK = RGBColor(0, 0, 0)

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)

st = doc.styles["Normal"]
st.font.name = KO
st.font.size = BODY
st.font.color.rgb = BLACK
st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)
st.paragraph_format.line_spacing = Pt(17)
st.paragraph_format.space_after = Pt(4)


def style_run(r, size=BODY, bold=False):
    r.font.name = KO
    r.font.size = size
    r.font.bold = bold
    r.font.color.rgb = BLACK
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.insert(0, rf)
    for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rf.set(qn(a), KO)
    return r


def p(text="", size=BODY, bold=False, after=4, before=0, indent=0,
      hang=None, align=None, line=17, keep=False):
    par = doc.add_paragraph()
    f = par.paragraph_format
    f.space_after, f.space_before, f.line_spacing = Pt(after), Pt(before), Pt(line)
    if indent:
        f.left_indent = Cm(indent)
    if hang is not None:
        f.first_line_indent = Cm(-hang)
    if align is not None:
        par.alignment = align
    f.keep_together = True
    if keep:
        f.keep_with_next = True
    if text:
        style_run(par.add_run(text), size, bold)
    return par


def rich(chunks, after=4, before=0, indent=0, hang=None, size=BODY):
    par = p("", after=after, before=before, indent=indent, hang=hang)
    for t, b in chunks:
        style_run(par.add_run(t), size, b)
    return par


def jang(text):
    return p(text, HEAD, True, after=8, before=18, keep=True)


def ne(text, bold=True):
    return p(f"□ {text}", BODY, bold, after=5, before=8, keep=True)


def o(text, after=4):
    return p(f"○ {text}", indent=0.65, hang=0.45, after=after)


def o_rich(chunks, after=4):
    return rich([("○ ", False)] + list(chunks), indent=0.65, hang=0.45, after=after)


def dash(text, after=3):
    return p(f"- {text}", SMALL, indent=1.25, hang=0.42, after=after)


def dash_rich(chunks, after=3):
    return rich([("- ", False)] + list(chunks), indent=1.25, hang=0.42,
                after=after, size=SMALL)


def table(rows, widths, small=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.find(qn("w:tblGrid")), widths):
        gc.set(qn("w:w"), str(w))
    sz = SMALL if small else BODY
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Pt(widths[ci] / 20)
            for li, line in enumerate(str(val).split("\n")):
                par = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.space_before = Pt(1)
                par.paragraph_format.line_spacing = Pt(14)
                par.paragraph_format.keep_with_next = (ri < len(rows) - 1)
                if ri == 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(par.add_run(line), sz, ri == 0)
        t.rows[ri]._tr.get_or_add_trPr().append(
            t.rows[ri]._tr.makeelement(qn("w:cantSplit"), {}))
    tblPr = t._tbl.tblPr
    b = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge, s in (("top", "6"), ("bottom", "6"), ("left", "4"), ("right", "4"),
                    ("insideH", "2"), ("insideV", "2")):
        e = b.makeelement(qn("w:" + edge), {})
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), s)
        e.set(qn("w:color"), "000000"); e.set(qn("w:space"), "0")
        b.append(e)
    tblPr.append(b)
    t.rows[0]._tr.get_or_add_trPr().append(
        t.rows[0]._tr.makeelement(qn("w:tblHeader"), {}))
    p("", after=8)
    return t


def rule_line(after=10):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(after)
    pPr = par._p.get_or_add_pPr()
    b = pPr.makeelement(qn("w:pBdr"), {})
    bot = b.makeelement(qn("w:bottom"), {})
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
    bot.set(qn("w:color"), "000000"); bot.set(qn("w:space"), "1")
    b.append(bot); pPr.append(b)


C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT

# ══ 표제부 ═════════════════════════════════════════════════════════════
p("AI 활용 정책 아이디어 평가체계 연구", SMALL, align=C, after=6)
p("연구 참여 계획서", TITLE, True, align=C, after=8)
p("— 참여 연구자로서 맡고자 하는 세 가지 역할 —", align=C, after=16)
p("국민대학교  김재준", align=R, after=2)
p("joyof15@gmail.com", SMALL, align=R, after=12)
rule_line(14)

ne("제안의 요지")
o("KDI 가 수행할 본 연구에 참여 연구자로 참가하고자 함")
o("맡을 수 있는 역할은 셋이며, 앞의 둘은 이미 만들어 둔 결과물이 근거가 "
  "되고 셋째는 이 연구에 이론 축을 더하는 일임")

table([
    ["", "역할", "근거 · 형태"],
    ["1", "평가체계 설계 자문",
     "작동하는 시스템을 설계·구현하고\n예비 검토까지 마친 경험"],
    ["2", "구현 — 일부 프로그래밍",
     "문답·채점 엔진, 선례 판정기,\n검증 코퍼스를 직접 제작"],
    ["3", "아이디어 내생 성장론 관점의 보고서",
     "정책제안 선별을 성장이론의\n문제로 다시 세우는 시론"],
], [500, 2800, 4700])

o_rich([("셋째가 이 제안의 핵심", True),
        ("임. 앞의 둘은 연구를 굴러가게 하는 일이고, 셋째는 이 연구가 왜 "
         "경제 연구기관에서 수행되어야 하는지를 답하는 일임", False)], after=8)

# ══ Ⅰ ══════════════════════════════════════════════════════════════════
jang("Ⅰ. 이 연구를 보는 관점")

ne("평가체계 연구로만 두면 도구 개선에 그침")
o("측정 신뢰도와 도구 격차는 그 자체로 중요한 주제이나, 평가 방법론의 "
  "문제로 한정하면 산출이 '더 나은 채점표'에 머무름")

ne("정책제안의 선별은 성장의 문제이기도 함")
o_rich([("내생 성장론은 아이디어를 성장의 동력으로 봄. 다만 그 논의는 "
         "아이디어의 ", False), ("생산", True),
        ("에 집중해 왔고, 생산된 아이디어 가운데 무엇을 실행할지 고르는 "
         "", False), ("선별", True),
        ("은 대체로 무비용·무오차로 가정되어 왔다고 봄", False)])
o("국민제안 제도는 아이디어 투입을 늘리는 장치이나, 선별 용량이 따라가지 "
  "못하면 투입 증가가 산출로 이어지지 않음")

ne("이 연구가 다루는 것이 바로 선별 용량임")
o("평가 시스템은 선별 비용을 낮추는 기술이고, 측정 재현성 문제는 선별이 "
  "잡음을 포함한다는 증거이며, 도구 격차는 선별 기준이 아이디어의 질에서 "
  "진술 능력으로 옮겨 가는 현상임")
o_rich([("세 가지 모두 ", False), ("성장의 언어로 다시 쓸 수 있다", True),
        ("고 보며, 그 작업을 맡고자 함", False)])

# ══ Ⅱ ══════════════════════════════════════════════════════════════════
jang("Ⅱ. 참여의 근거 — 이미 만들어 둔 것")

ne("작동하는 평가 시스템 일체")
o("소크라테스식 12문답으로 제안을 구조화하고, 국가 데이터 네 통로로 선례를 "
  "확인하여 이원 평가하는 웹 시스템")
dash("축 A(방어력) — 체크리스트 채점과 종합판단 채점을 분리. 둘의 괴리가 "
     "크면 낮은 쪽을 채택하는 규칙")
dash("축 B(실질 독창성) — 점수를 매기지 않고 선례 계열과 확신도로만 답함")
dash("설계·구현·시험을 모두 직접 수행하였으므로, 왜 그렇게 만들었고 어디가 "
     "약한지를 함께 설명할 수 있음")

ne("검증용 국가 데이터 코퍼스")
table([
    ["통로", "자료", "규모"],
    ["① 재정(집행)", "세출예산 사업", "14,122개 사업"],
    ["② 연구(검토)", "KDI 발간물", "7,362건"],
    ["③ 입법(시도)", "국회 의안", "열린국회정보 실시간 조회"],
    ["④ 해외(시행)", "OECD OPSI 정책사례", "1,015건 · 98개국"],
], [1600, 3500, 2900])

o_rich([("조회 결과를 ", False), ("미실행 / 조회 실패 / N건", True),
        (" 셋으로 구분하도록 설계함. '못 찾은 것'을 '없는 것'으로 읽는 오류를 "
         "자료 구조 수준에서 막은 것이며, 이 구분이 선별의 신뢰도를 논하는 "
         "근거가 됨", False)])

ne("예비 검토 결과 — 연구 질문의 출발점")
o_rich([("동일 로그를 4회 채점한 결과 합계가 ", False), ("13~17로 분포", True),
        ("함 (30점 척도)", False)])
o_rich([("AI 보조 조건의 점수가 높았으나 ", False), ("n=1쌍", True),
        ("의 통제되지 않은 비교로, 방향의 시사에 한정됨", False)])
o("격차의 대부분이 체크리스트에서 발생하고 종합판단에서는 거의 발생하지 "
  "않았음. 분량 가설은 지지되지 않았음")

ne("기관 서버에 설치 가능한 상태")
o("윈도우 서버판과 설치·운영 안내서를 정리하여 전달하였으며, 설치 즉시 "
  "다수 참여자 대상 자료 수집이 가능함")

# ══ Ⅲ ══════════════════════════════════════════════════════════════════
jang("Ⅲ. 맡고자 하는 역할")

ne("역할 1 — 평가체계 설계 자문")
o("문답 설계, 채점 기준, 선례 판정 규칙에 대한 설계 근거와 대안 제시")
o("실험 설계 단계에서 측정오차 추정 절차와 임계값 사전 결정 방식에 대한 의견 "
  "제시")
dash("무엇을 만들 수 있는가보다 무엇이 안 되는지를 아는 쪽이 자문의 값이라 "
     "보며, 예비 검토에서 스스로 확인한 실패가 그 자산임")

ne("역할 2 — 구현")
o("현행 시스템의 수정·확장. 실험 조건 분기, 반복채점 자동화, 자료 추출 "
  "도구 등 연구 수행에 필요한 부분")
o("KDI 연구진이 설계를 정하면 그것을 돌아가는 형태로 만드는 일을 맡고자 함")
dash("전체 개발을 맡겠다는 뜻이 아니며, 이미 있는 코드를 아는 사람이 손대는 "
     "편이 빠른 범위에 한정")

ne("역할 3 — 아이디어 내생 성장론 관점의 보고서")
o("본 연구의 실증 결과를 성장이론의 틀에서 해석하는 별도 보고서 집필")
o_rich([("다음 장에 구상을 적음. 이 역할이 ", False),
        ("앞의 둘과 성격이 다르므로", True),
        (" 지면을 나누어 설명함", False)])

# ══ Ⅳ ══════════════════════════════════════════════════════════════════
jang("Ⅳ. 세 번째 역할의 구상")

ne("출발 질문 — 아이디어는 정말 찾기 어려워졌는가")
o("연구생산성 하락, 곧 같은 투입으로 얻는 아이디어가 줄고 있다는 관찰은 "
  "성장 문헌의 오랜 논점임")
o_rich([("이를 뒤집어 물을 수 있다고 봄 — ", False),
        ("찾기 어려워진 것인가, 찾아 놓고 알아보지 못한 것인가", True)])

ne("아이디어의 비경합성은 '인지된 뒤'의 성질임")
o("아이디어가 비경합적이어서 한 번 발견되면 모두가 쓸 수 있다는 명제는, "
  "그 아이디어가 발견되고 인지된 다음에 성립함")
o_rich([("제안되었으나 선별되지 못한 아이디어는 ", False),
        ("사회적으로 존재하지 않는 것과 같음", True),
        (". 이 구간이 성장 모형에서 거의 다루어지지 않았다고 봄", False)])

ne("선별을 생산함수에 넣어 보는 시론")
o("아이디어 스톡의 증가를 '생산된 아이디어 × 선별 통과율'로 분해하고, "
  "선별 통과율을 검증 용량과 선별 정확도의 함수로 두는 단순한 확장을 "
  "시도하고자 함")
dash("검증 용량 — 담당 인력이 기간 내 처리할 수 있는 건수")
dash("선별 정확도 — 좋은 제안을 통과시키고 나쁜 제안을 거르는 정도. "
     "본 연구의 측정 재현성이 바로 이 값의 상한을 규정함")
o_rich([("이 틀에서 보면 ", False),
        ("측정 신뢰도는 방법론 문제가 아니라 성장의 파라미터", True),
        ("가 됨", False)])

ne("도구 격차를 배분의 왜곡으로 다시 씀")
o("진술 능력이 선별을 좌우하면, 선별 기준은 아이디어의 질이 아니라 진술 "
  "자본이 됨. 이는 아이디어 배분의 왜곡이며 사회적 손실로 계산될 수 있음")
o_rich([("생성형 AI 는 이 왜곡을 ", False), ("키울 수도 줄일 수도", True),
        (" 있음 — 금지하면 접근 가능한 소수에게만 남아 왜곡이 커지고, "
         "균등 제공하면 진술 자본의 차이가 상쇄되어 줄어듦", False)])
dash("탐지·금지 대 균등 제공의 선택이 윤리 문제가 아니라 배분 효율의 문제로 "
     "다시 놓임")

ne("정책 함의 — 투자가 한쪽에만 가 있음")
o("R&D 보조금은 아이디어 생산에 투자하나, 선별 용량에 대한 투자는 거의 "
  "논의되지 않음")
o("국민제안 제도가 그 예로, 접수는 대량이나 검증은 소수 담당자의 수작업에 "
  "의존함. 생산 측 투자만 늘리면 병목이 더 조여짐")
o_rich([("본 연구가 만드는 평가 시스템은 ", False),
        ("선별 용량에 대한 기술적 투자", True),
        ("의 한 형태이며, 그 효과를 성장의 언어로 계산해 보이는 것이 이 "
         "보고서의 목표임", False)])

ne("이 구상의 지위")
o_rich([("현 단계에서는 ", False), ("시론", True),
        ("임. 형식 모형의 타당성과 선행 문헌과의 관계는 문헌 검토를 거쳐야 "
         "확정할 수 있으며, 실증 결과가 나오기 전에는 수치를 넣지 않을 것임",
         False)])
o("본 연구의 실증이 부정적으로 나오더라도 — 격차가 확인되지 않거나 "
  "재현성이 충분한 것으로 밝혀지더라도 — 그 결과를 그대로 반영하여 쓸 것임")

# ══ Ⅴ ══════════════════════════════════════════════════════════════════
jang("Ⅴ. 참여 방식")

ne("주관은 KDI, 본인은 참여자")
o("연구의 방향과 설계 결정은 KDI 연구진이 주관하고, 본인은 위 세 역할의 "
  "범위에서 참여하고자 함")
o("참여 형태(공동연구·위탁·자문 등)와 범위는 기관 기준에 따르겠음")

ne("자료와 도구는 연구에 넘김")
o("평가 시스템, 검증 코퍼스, 채점 기준, 예비 검토 원자료를 본 연구에 제공")
o_rich([("연구 종료 후에도 기관 자산으로 남도록 하고자 함. ", False),
        ("이 연구가 끝나면 쓸 수 없는 도구가 되지 않게 하는 것", True),
        ("이 참여의 조건 가운데 하나임", False)])

ne("역할 1·2 와 역할 3 을 분리")
o("자문·구현은 연구 수행 기간에 걸쳐 수시로, 보고서는 실증 결과가 나온 "
  "뒤에 착수하는 것이 순서에 맞다고 봄")

# ══ Ⅵ ══════════════════════════════════════════════════════════════════
jang("Ⅵ. 밝혀 둘 것")

ne("도구를 만든 사람이 그 도구를 평가하는 연구에 참여함")
o("설계를 옹호하는 방향의 해석이 끼어들 수 있는 위치임")
o_rich([("채점 기준과 예비 검토 원자료를 공개하여 제3자 재분석이 가능하게 "
         "하고, ", False),
        ("채점 결과의 해석은 본인이 아닌 연구진이 맡는 것", True),
        ("이 타당하다고 봄", False)])

ne("예비 관찰은 확정된 결과가 아님")
o_rich([("특히 AI 보조 조건의 점수 차이는 ", False), ("n=1쌍", True),
        ("의 통제되지 않은 비교임. 본실험에서 재현되지 않을 가능성이 있으며 "
         "그 경우 그대로 보고되어야 함", False)])

ne("구조적 순환이 남아 있음")
o("문답 생성·채점·선례 판정에 같은 계열 모델이 쓰임. 모델이 잘 만드는 답을 "
  "모델이 높게 치는 구조일 수 있으며, 전문가 심사를 준거로 병행하는 것이 "
  "완화책이나 완전한 해소는 아님")

ne("역할 3 의 이론 작업은 검증을 거쳐야 함")
o("내생 성장론 문헌에서 선별·심사 비용을 다룬 선행 연구가 이미 있을 수 "
  "있음. 문헌 검토를 선행하여 신규성 주장의 수위를 정하겠음")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "연구참여계획서_KDI.docx")
doc.save(out)
print("저장 완료 —", out)
