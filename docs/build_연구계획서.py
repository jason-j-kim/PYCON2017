#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""연구계획서를 만든다 — 개인이 앞으로 할 연구를 밝히는 문서.

기관의 사업계획이 아니다. 그래서 일정표·투입표·소요예산·행정사항을 넣지
않는다. 대신 무엇을 밝히려 하는지, 왜 그 질문인지, 어떻게 확인할 것인지,
무엇을 스스로 경계하는지를 적는다.

서식은 개조식(□ ○ - ·)을 유지한다. 읽는 사람이 그 형식에 익숙하기
때문이고, 주장과 근거의 층위가 눈에 바로 보이기 때문이다.

숫자는 docs/KDI연구초안.md 의 예비 검토 결과에서만 가져온다. 예비 관찰은
예비 관찰로 적는다.
"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

KO = "맑은 고딕"
TITLE, HEAD, BODY, SMALL = Pt(15), Pt(12), Pt(11), Pt(10)
BLACK = RGBColor(0, 0, 0)

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)

st = doc.styles["Normal"]
st.font.name = KO
st.font.size = BODY
st.font.color.rgb = BLACK
st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)
st.paragraph_format.line_spacing = Pt(17)
st.paragraph_format.space_after = Pt(4)


def style_run(r, size=BODY, bold=False):
    r.font.name = KO
    r.font.size = size
    r.font.bold = bold
    r.font.color.rgb = BLACK
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.insert(0, rf)
    for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rf.set(qn(a), KO)
    return r


def p(text="", size=BODY, bold=False, after=4, before=0, indent=0,
      hang=None, align=None, line=17, keep=False):
    par = doc.add_paragraph()
    f = par.paragraph_format
    f.space_after, f.space_before, f.line_spacing = Pt(after), Pt(before), Pt(line)
    if indent:
        f.left_indent = Cm(indent)
    if hang is not None:
        f.first_line_indent = Cm(-hang)
    if align is not None:
        par.alignment = align
    f.keep_together = True
    if keep:
        f.keep_with_next = True
    if text:
        style_run(par.add_run(text), size, bold)
    return par


def rich(chunks, after=4, before=0, indent=0, hang=None, size=BODY):
    par = p("", after=after, before=before, indent=indent, hang=hang)
    for t, b in chunks:
        style_run(par.add_run(t), size, b)
    return par


def jang(text):
    return p(text, HEAD, True, after=8, before=18, keep=True)


def ne(text, bold=True):
    return p(f"□ {text}", BODY, bold, after=5, before=8, keep=True)


def o(text, after=4):
    return p(f"○ {text}", indent=0.65, hang=0.45, after=after)


def o_rich(chunks, after=4):
    return rich([("○ ", False)] + list(chunks), indent=0.65, hang=0.45, after=after)


def dash(text, after=3):
    return p(f"- {text}", SMALL, indent=1.25, hang=0.42, after=after)


def dash_rich(chunks, after=3):
    return rich([("- ", False)] + list(chunks), indent=1.25, hang=0.42,
                after=after, size=SMALL)


def table(rows, widths, small=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.find(qn("w:tblGrid")), widths):
        gc.set(qn("w:w"), str(w))
    sz = SMALL if small else BODY
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Pt(widths[ci] / 20)
            for li, line in enumerate(str(val).split("\n")):
                par = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.space_before = Pt(1)
                par.paragraph_format.line_spacing = Pt(14)
                par.paragraph_format.keep_with_next = (ri < len(rows) - 1)
                if ri == 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(par.add_run(line), sz, ri == 0)
        t.rows[ri]._tr.get_or_add_trPr().append(
            t.rows[ri]._tr.makeelement(qn("w:cantSplit"), {}))
    tblPr = t._tbl.tblPr
    b = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge, s in (("top", "6"), ("bottom", "6"), ("left", "4"), ("right", "4"),
                    ("insideH", "2"), ("insideV", "2")):
        e = b.makeelement(qn("w:" + edge), {})
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), s)
        e.set(qn("w:color"), "000000"); e.set(qn("w:space"), "0")
        b.append(e)
    tblPr.append(b)
    t.rows[0]._tr.get_or_add_trPr().append(
        t.rows[0]._tr.makeelement(qn("w:tblHeader"), {}))
    p("", after=8)
    return t


def rule_line(after=10):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(after)
    pPr = par._p.get_or_add_pPr()
    b = pPr.makeelement(qn("w:pBdr"), {})
    bot = b.makeelement(qn("w:bottom"), {})
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
    bot.set(qn("w:color"), "000000"); bot.set(qn("w:space"), "1")
    b.append(bot); pPr.append(b)


C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT

# ══ 표제부 ═════════════════════════════════════════════════════════════
p("생성형 AI 시대의 공공 평가", SMALL, align=C, after=6)
p("정책 아이디어 평가체계에 관한 연구계획", TITLE, True, align=C, after=8)
p("— 평가도구의 재현성과 도구 격차를 중심으로 —", align=C, after=16)
p("국민대학교  김재준", align=R, after=2)
p("joyof15@gmail.com", SMALL, align=R, after=12)
rule_line(14)

ne("연구의 목적")
o("서면 진술에 의존하는 공공 평가가 생성형 AI 확산으로 무엇을 재고 있는지 "
  "불분명해진 상황에서, 그 영향의 크기를 실증하고 대응 원칙을 제시하고자 함")
o_rich([("작동하는 평가 시스템을 직접 만들어 두었으므로, 사변이 아니라 "
         "", False),
        ("실측으로 답하는 것", True), ("을 이 연구의 성격으로 삼고자 함", False)],
       after=8)

# ══ Ⅰ ══════════════════════════════════════════════════════════════════
jang("Ⅰ. 문제의 소재")

ne("공공 평가는 제안을 보는 것이 아니라 제안서를 봄")
o("국민제안·공모사업·연구과제 선정에서 심사 대상은 착상 자체가 아니라 "
  "신청자가 작성한 문서임")
o("이 구조는 진술 능력과 착상의 질이 대체로 함께 간다는 암묵적 전제 위에 "
  "서 있었음")

ne("생성형 AI 가 그 전제를 무너뜨림")
o("같은 착상이라도 AI 보조 여부에 따라 진술의 완성도가 크게 달라짐")
o_rich([("이 경우 평가는 제안의 질이 아니라 ", False),
        ("도구 접근성", True),
        ("을 재게 되며, 이는 측정하려던 것과 다른 것을 측정하는 문제임", False)])

ne("탐지·금지로는 풀리지 않는다고 봄")
o("AI 생성물 탐지의 정확도가 불이익 처분의 근거가 될 수준에 이르지 못함")
o_rich([("평가 대상이 '사람의 능력'이 아니라 '제안이라는 산출물'인 영역에서는 "
         "금지보다 ", False), ("균등 제공", True),
        ("이 정합적일 수 있다고 보나, 이는 격차의 존재가 확인되어야 성립하는 "
         "주장이므로 본 연구의 검증 대상으로 둠", False)])

ne("연구자로서 이 문제에 접근하는 방식")
o("평가 제도를 바깥에서 논평하는 대신, 대안이 될 만한 평가 절차를 실제로 "
  "만들어 돌려 보고 그 과정에서 나오는 것을 자료로 삼고자 함")

# ══ Ⅱ ══════════════════════════════════════════════════════════════════
jang("Ⅱ. 지금까지 한 일")

ne("평가 시스템을 만들어 작동시킴")
o("정책 아이디어를 소크라테스식 12문답으로 구조화하고, 국가 데이터 네 개 "
  "통로로 선례를 확인하여 이원 평가하는 웹 시스템")
dash("축 A(방어력) — 독창성·실용성·수용태도 3기준. 체크리스트 채점과 "
     "종합판단 채점을 따로 두고, 둘의 괴리가 크면 낮은 쪽을 채택")
dash("축 B(실질 독창성) — 네 통로 조회 후 선례 계열과 확신도를 판정. "
     "점수를 매기지 않고 계열로만 답함")

ne("검증에 쓸 국가 데이터를 구축·탑재")
table([
    ["통로", "자료", "규모"],
    ["① 재정(집행)", "세출예산 사업", "14,122개 사업"],
    ["② 연구(검토)", "KDI 발간물", "7,362건"],
    ["③ 입법(시도)", "국회 의안", "열린국회정보 실시간 조회"],
    ["④ 해외(시행)", "OECD OPSI 정책사례", "1,015건 · 98개국"],
], [1600, 3500, 2900])

o_rich([("조회 결과를 ", False), ("미실행 / 조회 실패 / N건", True),
        (" 셋으로 구분하도록 설계함. 앞의 둘은 '선례가 없다'가 아니라 "
         "'확인하지 못했다'는 뜻이며, 판정 근거로 쓰지 않고 확신도 상향을 "
         "막음", False)])
dash("평가에서 가장 흔한 오류가 '못 찾은 것'을 '없는 것'으로 읽는 것이라 "
     "보아, 이를 자료 구조 수준에서 분리함")

ne("연구 자료가 쌓이는 상태로 배포 가능하게 준비")
o("문답·채점·선례 판정이 전부 기록되며, 건별로 문답 전문과 평가보고서를 "
  "내려받을 수 있음")
o("기관 서버에 설치할 수 있는 형태로 정리하여, 다수 참여자를 대상으로 한 "
  "자료 수집이 가능한 상태임")

# ══ Ⅲ ══════════════════════════════════════════════════════════════════
jang("Ⅲ. 예비 관찰과 그 함의")

ne("자체 검토에서 네 가지가 관찰됨 — 모두 예비 수준임")

o_rich([("(관찰 1) ", True),
        ("동일한 대화 로그를 체크리스트로 4회 채점한 결과 합계가 ", False),
        ("13~17로 분포", True), ("함 (30점 척도, 평균 15.0)", False)])
dash("체크리스트는 '해당 사건이 대화에 있었는가'만 판정하는 기계적 절차로 "
     "설계하였고 재현성이 그 설계의 근거였으나, 경험적으로 성립하지 않음")

o_rich([("(관찰 2) ", True),
        ("동일 아이디어에 대해 AI 보조 조건의 점수가 높았음 ", False),
        ("(5.2 → 8.0)", True)])
dash_rich([("다만 ", False), ("n=1쌍", True),
           ("의 통제되지 않은 비교이며 두 조건에 서로 다른 질문이 제시되었음. "
            "방향의 시사에 한정되고 크기의 근거로 쓸 수 없음", False)])

o_rich([("(관찰 3) ", True),
        ("격차의 발생 지점을 분해한 결과 대부분이 ", False),
        ("체크리스트에서 발생", True),
        ("하였고, 내용의 질을 보는 종합판단에서는 거의 발생하지 않았음", False)])
dash("관측된 체크리스트 차이(14점)는 위 측정 변동폭(4점)의 3.5배로 잡음만으로는 "
     "설명되지 않으나, 조건 비교 자체가 통제되지 않아 잠정적임")

o_rich([("(관찰 4) ", True),
        ("분량 가설은 지지되지 않았음", False)])
dash("동일 내용을 1.46배 길이로 재작성한 로그의 체크리스트 점수가 오르지 "
     "않음 (평균 15.0 → 13.0). 격차의 원인이 분량이 아니라 채점 항목의 "
     "열거일 가능성")

ne("이 관찰들이 연구의 순서를 바꿈")
o_rich([("당초에는 '도구 격차를 측정한다'가 연구 질문이었으나, ", False),
        ("측정도구 자체가 재현되지 않는다면 그 질문은 성립하지 않음", True)])
o("도구의 측정오차를 모르는 채 격차를 보고하면, 그 수치가 실재하는 차이인지 "
  "채점 잡음인지 판별할 수 없기 때문임")
o_rich([("따라서 ", False), ("재현성 확인을 격차 측정에 선행", True),
        ("시키는 것으로 설계를 바꾸고자 함", False)])

# ══ Ⅳ ══════════════════════════════════════════════════════════════════
jang("Ⅳ. 앞으로 밝히려는 것")

ne("세 질문을 순서대로 다루고자 함")

table([
    ["", "질문", "성격"],
    ["질문 1", "이 평가도구는 재현되는가\n— 동일 입력에 동일 판정이 나오는가",
     "선결 과제"],
    ["질문 2", "생성형 AI 보조가 평가 결과를 얼마나 바꾸는가\n"
               "— 그 차이는 측정오차를 넘는가", "실증"],
    ["질문 3", "그 결과는 국민제안 처리 제도에 무엇을 뜻하는가",
     "제도 설계"],
], [900, 5200, 1900])

o_rich([("질문 1이 부정되면 질문 2로 넘어가지 않음", True),
        (". 재현되지 않는 도구로 격차를 재는 것은 순서가 뒤바뀐 일이며, "
         "그 경우 연구는 '평가도구를 어떻게 재현 가능하게 만들 것인가'로 "
         "방향을 바꾸게 됨", False)])

o_rich([("질문 3은 질문 2의 결과에 의존함", True),
        (". 격차가 확인되지 않으면 제도 개편을 권고하지 않을 것임. "
         "확인되지 않은 것을 전제로 제도를 설계하면 방어할 수 없음", False)])

# ══ Ⅴ ══════════════════════════════════════════════════════════════════
jang("Ⅴ. 어떻게 확인할 것인가")

ne("측정오차를 먼저 추정하고 임계값을 사전에 정함")
o("동일 로그를 반복 채점하여 도구의 변동폭을 구하고, '이 크기를 넘어야 "
  "격차로 본다'는 기준을 자료를 보기 전에 확정하고자 함")
dash("사후에 기준을 정하면 원하는 결론이 나오는 쪽으로 기울게 됨")

ne("같은 사람이 같은 아이디어를 두 조건에서 진술하게 함")
o("예비 관찰의 결정적 약점이 두 조건에 서로 다른 질문이 제시된 점이었으므로, "
  "질문을 고정하고 참여자 내 대응 비교로 설계하고자 함")

ne("격차가 어디서 생기는지 분해함")
o("체크리스트와 종합판단을 따로 집계하여, 차이가 내용의 질에서 오는지 "
  "진술의 열거 방식에서 오는지 가름")
dash("관찰 3이 옳다면 격차는 '무엇을 말했는가'가 아니라 '빠짐없이 열거했는가'"
     "에서 생기는 것이며, 이는 채점 설계를 고쳐 줄일 수 있는 종류의 격차임")

ne("사람의 심사를 준거로 병행함")
o("문답 생성·채점·판정에 같은 계열 모델이 쓰이는 구조적 순환이 있으므로, "
  "동일 자료를 정책 전문가가 독립적으로 심사하게 하여 대조하고자 함")

ne("완화 설계를 만들어 시험함")
o("격차가 확인될 경우, 채점 항목을 '언급 여부'에서 '검증 가능성'으로 옮기는 "
  "등의 설계 변경이 격차를 줄이는지 검증")
dash("현재의 완화 설계는 같은 자료로 고안한 것이므로 표본 내 적합에 불과함. "
     "별도 자료로 확인하는 것을 과제로 둠")

ne("자료 수집을 겹쳐 씀")
o("질문 1·2를 위한 세션에 전문가 심사를 함께 붙이면 질문 3의 선결 과제인 "
  "준거 타당도가 동시에 확보됨. 별도 수집을 벌이지 않고자 함")

# ══ Ⅵ ══════════════════════════════════════════════════════════════════
jang("Ⅵ. 내놓으려는 것")

ne("논문")
o("(실증) 공공 평가도구의 재현성과 생성형 AI 도구 격차에 관한 실증 분석")
o("(제도) 국민 정책제안의 검증 병목과 처리 절차 재설계안")

ne("도구와 자료")
o("평가 시스템 일체 — 문답 설계, 채점 기준 30항목, 선례 판정 규칙")
o("검증용 코퍼스 구축 방법과 커버리지 명세")
o_rich([("반복채점 원자료와 전사 자료를 함께 공개하고자 함", True),
        (". 재현성을 묻는 연구가 스스로 재현 가능하지 않으면 앞뒤가 맞지 "
         "않음", False)])

ne("실무에 바로 쓰이는 형태")
o("설치하면 그대로 돌아가는 배포본과 설치·운영 안내서")
o("이 연구의 산출이 논문에 그치지 않고 남는 도구가 되도록 하고자 함")

# ══ Ⅶ ══════════════════════════════════════════════════════════════════
jang("Ⅶ. 스스로 경계하는 것")

ne("예비 관찰을 결과처럼 쓰지 않을 것")
o_rich([("관찰 2는 ", False), ("n=1쌍", True),
        ("의 통제되지 않은 비교임. 본실험에서 재현되지 않을 가능성을 배제하지 "
         "않으며, 그 경우 그대로 보고할 것임", False)])

ne("순환 구조를 인정하고 다룰 것")
o("문항을 만들고 답을 채점하고 선례를 판정하는 일에 같은 계열 모델이 쓰임. "
  "모델이 잘 만드는 답을 모델이 높게 치는 구조일 수 있음")
o("사람의 심사를 준거로 두는 것이 완화책이나 완전한 해소는 아니며, 이 한계를 "
  "논문에 명시할 것임")

ne("도구를 만든 사람이 그 도구를 평가한다는 위치를 밝힐 것")
o("설계자가 곧 연구자인 상황에서는 설계를 옹호하는 방향의 해석이 끼어들기 "
  "쉬움. 채점 기준과 원자료를 공개하여 제3자 재분석이 가능하게 하는 것으로 "
  "대응하고자 함")

ne("일반화의 범위를 좁게 잡을 것")
o("정책 아이디어라는 한 종류의 산출물, 한 평가 절차에서 얻은 결과임. "
  "공공 평가 전반으로 넓히는 주장은 하지 않을 것임")

ne("사람을 대상으로 하는 자료 수집의 절차를 지킬 것")
o("참여자 동의와 연구윤리 심의를 선행하며, 문답 기록에 개인정보가 포함될 수 "
  "있으므로 수집·보관·파기 절차를 사전에 정하고자 함")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "연구계획서_정책아이디어평가.docx")
doc.save(out)
print("저장 완료 —", out)
